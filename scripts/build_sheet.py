#!/usr/bin/env python3
"""Build sheets/LLM-Roadmap.xlsx (imported into Google Sheets) from the repo.

  .venv/bin/python scripts/build_sheet.py

Reads curriculum.json, projects.json, schedule.csv, progress.csv and outlines/*.json.
Re-run after outlines change, then re-import (File > Import > Replace spreadsheet).
"""
import csv, datetime as dt, glob, json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
C = lambda p: f"{ROOT}/curriculum/{p}"
cur = json.load(open(C("curriculum.json")))
projects = json.load(open(C("projects.json")))
sched = list(csv.DictReader(open(C("schedule.csv"))))
prog = {r["code"]: r for r in csv.DictReader(open(C("progress.csv")))}
outl = {}
for f in sorted(glob.glob(C("outlines/M*.json"))):
    try:
        for o in json.load(open(f)):
            outl[o["code"]] = o
    except Exception as ex:
        print("skip", f, ex)
mods = {m["code"]: m for m in cur["modules"]}
eps = cur["episodes"]
src = cur["sources"]

INK, PAPER, MUTED = "111827", "FFFFFF", "6B7280"
HEAD = PatternFill("solid", fgColor="111827")
SUB = PatternFill("solid", fgColor="F3F4F6")
thin = Side(style="thin", color="E5E7EB")
BORDER = Border(bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
wb = Workbook()


def sheet(title, headers, widths, rows, freeze="B2", wrap=True, height=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color=PAPER); c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if wrap:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP
    if height:
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = height
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 30
    return ws


def yn(ws, col_letters, n):
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    green = PatternFill("solid", fgColor="D1FAE5")
    for L in col_letters:
        dv.add(f"{L}2:{L}{n + 1}")
        ws.conditional_formatting.add(f"{L}2:{L}{n + 1}",
                                      CellIsRule(operator="equal", formula=['"Y"'], fill=green))


def mod_fill(code):
    return PatternFill("solid", fgColor=mods[code]["color"].lstrip("#"))


# ------------------------------------------------------------------ Start Here
ws = wb.active; ws.title = "Start Here"
first_proj = {}
for r in sched:
    if r["kind"] == "project":
        first_proj.setdefault(r["code"], r["date"])
last_proj = {}
for r in sched:
    if r["kind"] == "project":
        last_proj[r["code"]] = r["date"]
start_txt = [
    [cur["series"] + " — the roadmap"],
    [f"{len(eps)} episodes · {len(cur['modules'])} modules · {len(projects)} flagship projects · "
     f"{sched[0]['date']} → {sched[-1]['date']}"],
    [""],
    ["How to use this sheet"],
    ["1. Today tab: what to record today (auto-filters the Calendar by today's date)."],
    ["2. Night before: open Episodes & Scripts, read tomorrow's two rows (hook + beats). The full prep sheet is in the repo at the 'Prep sheet' path."],
    ["3. Record in Meld (TALK/CODE scenes). Then tick Recorded on the Calendar row; Claude runs the HyperFrames edit + Short + explainer."],
    ["4. Review renders, upload, tick each platform column. Sunday: fill the Scorecard."],
    ["5. Flagship projects (Projects tab) are the resume. Episodes are the path to them. On sprint days, slot B is a filmed project build log."],
    [""],
    ["The rule of the series"],
    ["Every episode builds or measures something real and ends with a RECEIPT: allclose against the production implementation, a gradcheck, a benchmark vs the hardware peak, or a published number."],
    [""],
    ["Repo: ~/Documents/DL-content   ·   CLI: dl today | dl start | dl prep | dl explainer | dl done   ·   Sources of truth: curriculum.json (content), this sheet (daily status)"],
    [""],
    ["Milestones"],
]
for p in projects:
    if p["id"] in first_proj:
        start_txt.append([f"{p['id']} {p['name']}: sprint {first_proj[p['id']]} → {last_proj[p['id']]}"])
start_txt += [[""], ["Apply windows (see Roadmap tab for the reasoning)"],
              ["• Now: applied AI / agent infra / iOS-at-labs roles (existing OperatorOK record)."],
              [f"• After P5 + P1 (~{last_proj.get('P1', '')}): inference, kernels and ML-infra roles; Anthropic Fellows application with P5/P1 + the P7 plan."],
              [f"• After P2 + P3 (~{last_proj.get('P3', '')}): lab research-engineer roles (inference, pretraining infra)."],
              [f"• After P4 + P6 + P7 (~{last_proj.get('P7', '')}): post-training / RL / safety research-engineer roles at every frontier lab."]]
for r in start_txt:
    ws.append(r)
ws.column_dimensions["A"].width = 150
ws["A1"].font = Font(bold=True, size=18)
ws["A2"].font = Font(color=MUTED)
for r in range(1, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value or ""
    if v in ("How to use this sheet", "The rule of the series", "Milestones") or v.startswith("Apply windows"):
        ws.cell(row=r, column=1).font = Font(bold=True, size=13)
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True)

# ------------------------------------------------------------------ Today
ws = wb.create_sheet("Today")
ws["A1"] = "Today"; ws["A1"].font = Font(bold=True, size=16)
ws["B1"] = "=TODAY()"; ws["B1"].number_format = "dddd, mmm d yyyy"
ws["A3"] = "Recording today"; ws["A3"].font = Font(bold=True)
ws["A4"] = '=IFERROR(FILTER(Calendar!B2:E, Calendar!A2:A=TODAY()), "Nothing scheduled today")'
ws["A8"] = "Tomorrow (read the prep sheets tonight)"; ws["A8"].font = Font(bold=True)
ws["A9"] = '=IFERROR(FILTER(Calendar!B2:E, Calendar!A2:A=TODAY()+1), "Nothing scheduled")'
ws["A13"] = "Scripts for today"; ws["A13"].font = Font(bold=True)
ws["A14"] = ('=IFERROR(FILTER({\'Episodes & Scripts\'!B2:B, \'Episodes & Scripts\'!F2:F, \'Episodes & Scripts\'!G2:G}, '
             '\'Episodes & Scripts\'!D2:D=TODAY()), "")')
ws["A20"] = "Progress"; ws["A20"].font = Font(bold=True)
ws["A21"] = "Episodes recorded"; ws["B21"] = "=COUNTIF('Episodes & Scripts'!P2:P,\"Y\")"
ws["A22"] = "Episodes published (YouTube)"; ws["B22"] = "=COUNTIF('Episodes & Scripts'!R2:R,\"Y\")"
ws["A23"] = "Total episodes"; ws["B23"] = len(eps)
ws["A24"] = "Behind / ahead (episodes)"; ws["B24"] = "=B21-COUNTIFS('Episodes & Scripts'!D2:D,\"<\"&TODAY())"
for col, w in zip("ABCDE", (30, 22, 60, 60, 20)):
    ws.column_dimensions[col].width = w

# ------------------------------------------------------------------ Roadmap
rm = []
dates = {}
for r in sched:
    if r["kind"] == "episode":
        m = r["code"][:3]
        dates.setdefault(m, [r["date"], r["date"]])[1] = r["date"]
proj_by_mod = {}
for p in projects:
    for m in [x.strip() for x in p["unlocked_by"].replace("ongoing from", "").split(",")]:
        proj_by_mod.setdefault(m, []).append(p["id"])
phase = {"M00": "1 Foundations", "M01": "1 Foundations", "M02": "1 Foundations", "M03": "1 Foundations",
         "M04": "2 Build the LLM", "M05": "2 Build the LLM", "M06": "2 Build the LLM",
         "M07": "3 Systems", "M08": "3 Systems", "M09": "3 Systems", "M10": "3 Systems",
         "M11": "4 Post-training + serving", "M12": "4 Post-training + serving", "M13": "4 Post-training + serving",
         "M14": "5 Research", "M15": "5 Research", "M16": "5 Research", "M17": "5 Research"}
for m in cur["modules"]:
    keys = []
    for e in eps:
        if e["module"] == m["code"]:
            for s in e["sources"]:
                if s not in keys:
                    keys.append(s)
    rm.append([phase[m["code"]], m["code"], m["name"], m["episodes"], dates[m["code"]][0], dates[m["code"]][1],
               m["hardware"], ", ".join(proj_by_mod.get(m["code"], [])), m["resume"],
               "\n".join(f"{src[k]['name']}" for k in keys[:6])])
ws = sheet("Roadmap", ["Phase", "Module", "Name", "Eps", "Start", "End", "Hardware", "Feeds project",
                       "Resume line it earns (fill X/Y with real numbers)", "Primary sources"],
           [22, 8, 34, 6, 12, 12, 24, 12, 60, 60], rm, freeze="D2")
for i, m in enumerate(cur["modules"], 2):
    ws.cell(row=i, column=2).fill = mod_fill(m["code"])
    ws.cell(row=i, column=2).font = Font(bold=True, color=PAPER)

# ------------------------------------------------------------------ Calendar
rows = []
for r in sched:
    d = dt.date.fromisoformat(r["date"])
    rows.append([d, r["dow"], r["slot"], r["code"], r["title"], r["kind"], "N" if r["kind"] != "review" else "", ""])
ws = sheet("Calendar", ["Date", "Day", "Slot", "Code", "What", "Kind", "Done", "Notes"],
           [12, 6, 6, 9, 70, 10, 7, 40], rows, freeze="A2", wrap=False)
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).number_format = "yyyy-mm-dd ddd"
    kind = ws.cell(row=r, column=6).value
    code = ws.cell(row=r, column=4).value or ""
    if kind == "review":
        for c in range(1, 9):
            ws.cell(row=r, column=c).fill = SUB
    elif kind == "project":
        ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor="FDE68A")
    elif code[:3] in mods:
        ws.cell(row=r, column=4).fill = mod_fill(code[:3]); ws.cell(row=r, column=4).font = Font(color=PAPER, bold=True)
yn(ws, ["G"], ws.max_row - 1)
ws.conditional_formatting.add(f"A2:H{ws.max_row}", FormulaRule(formula=["$A2=TODAY()"],
                              fill=PatternFill("solid", fgColor="DBEAFE")))
ws.auto_filter.ref = f"A1:H{ws.max_row}"

# ------------------------------------------------------------------ Episodes & Scripts
hdr = ["Ep", "Code", "Module", "Planned", "Title", "Hook (say first)", "Script beats (timed)", "Key math",
       "Code task", "Receipt (the proof)", "Gotchas", "Interview Q + answer", "Explainer narration (Manim)",
       "Short (vertical cut)", "Prep sheet ready", "Recorded", "Edited", "YouTube", "Short posted", "X",
       "LinkedIn", "Instagram", "TikTok", "Facebook", "YouTube URL", "YouTube title", "Thumbnail", "X post",
       "LinkedIn post", "Hardware", "Learn", "Builds / verifies", "Sources", "Prep sheet path", "Code file",
       "Explainer scene", "Prereqs", "Notes"]
rows = []
for e in eps:
    o, p = outl.get(e["code"], {}), prog.get(e["code"], {})
    g = lambda k: p.get(k, "N") or "N"
    rows.append([e["ep"], e["code"], f"{e['module']} {e['module_name']}",
                 dt.date.fromisoformat(p["planned_date"]) if p.get("planned_date") else "",
                 e["title"], o.get("hook", ""), "\n".join(o.get("beats", [])), o.get("key_math", ""),
                 o.get("code_task", ""), o.get("receipt", ""), o.get("gotchas", ""), o.get("interview_q", ""),
                 "\n".join(o.get("explainer", [])), o.get("short", ""),
                 g("prep_ready"), g("recorded"), g("edited"), g("yt_long"), g("short"), g("x"), g("linkedin"),
                 g("instagram"), g("tiktok"), g("facebook"), p.get("youtube_url", ""),
                 o.get("yt_title", ""), o.get("thumbnail", ""), o.get("x_post", ""), o.get("linkedin", ""),
                 e["hardware"], e["learn"], e["build"],
                 "\n".join(f"{src[k]['name']} — {src[k]['url']}" for k in e["sources"]),
                 e["prep"], e["code_file"], e["explainer"], o.get("prereq", ""), p.get("notes", "")])
ws = sheet("Episodes & Scripts", hdr,
           [5, 9, 18, 11, 34, 40, 80, 40, 40, 40, 36, 44, 60, 36] + [8] * 10 + [24, 40, 18, 44, 44, 18, 36, 36, 50, 30, 30, 30, 14, 30],
           rows, freeze="F2", height=150)
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
    code = ws.cell(row=r, column=2).value
    ws.cell(row=r, column=2).fill = mod_fill(code[:3]); ws.cell(row=r, column=2).font = Font(color=PAPER, bold=True)
yn(ws, [get_column_letter(i) for i in range(15, 25)], len(eps))
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{ws.max_row}"

# ------------------------------------------------------------------ Projects
rows = [[p["id"], p["name"], p["unlocked_by"], p["start_after"], first_proj.get(p["id"], "ongoing"),
         last_proj.get(p["id"], ""), p["weeks"], p["bar"], p["metric"], p["deliverable"], p["roles"],
         p["upstream"], "Not started", ""] for p in projects]
ws = sheet("Projects", ["ID", "Project", "Unlocked by", "Starts after", "Sprint start", "Sprint end", "Weeks",
                        "The bar (non-negotiable)", "Metric", "Deliverable", "Roles it targets",
                        "Upstream contribution", "Status", "Repo / write-up URL"],
           [5, 36, 12, 10, 12, 12, 6, 60, 40, 44, 40, 40, 14, 30], rows, freeze="C2", height=120)
dv = DataValidation(type="list", formula1='"Not started,Building,Benchmarking,Writing up,Shipped"')
ws.add_data_validation(dv); dv.add(f"M2:M{ws.max_row}")

# ------------------------------------------------------------------ DeepSeek Watch
ds_keys = [k for k in src if k.startswith("DS") or k in ("DEEPEP", "NSA", "R1", "GRPO", "DSV2", "DSV3")]
rows = []
for k in ds_keys:
    used = [e["code"] for e in eps if k in e["sources"]]
    rows.append([src[k]["name"], src[k]["kind"], src[k]["url"], ", ".join(used), "To read", "", ""])
rows.append(["(add new DeepSeek releases here as they ship: check github.com/deepseek-ai and huggingface.co/deepseek-ai weekly)", "", "https://github.com/deepseek-ai", "", "", "", ""])
ws = sheet("DeepSeek Watch", ["Paper / repo", "Type", "Link", "Episodes", "Status", "Key idea in one line", "Reproduced? (link)"],
           [60, 12, 60, 24, 12, 60, 30], rows, freeze="B2")
dv = DataValidation(type="list", formula1='"To read,Read,Implemented,Reproduced"')
ws.add_data_validation(dv); dv.add(f"E2:E{ws.max_row}")

# ------------------------------------------------------------------ Timetable
tt = [
    ["06:45–07:00", "Setup", "dl today · camera, mic, nvim :RecordMode, Meld scenes", "—"],
    ["07:00–07:45", "DSA drill", "dsa due: re-solve from blank file, timed", "Reps logged"],
    ["07:45–09:15", "DL episode A", "Record slot A (TALK/CODE). Mistakes stay in. Then: dl done <code> → Claude runs the edit", "Raw video A"],
    ["09:15–10:30", "DSA recording", "2 DSA episodes (dsa next / dsa start)", "2 raw DSA videos"],
    ["10:30–12:00", "Applications", "5 tailored applications + 3 outreach messages (job-search/tracker.csv)", "Tracker rows → Applied"],
    ["12:00–13:00", "Lunch + walk/gym", "Not optional", ""],
    ["13:00–14:30", "DL episode B / build log", "Slot B: next episode, or the project build log on sprint days", "Raw video B"],
    ["14:30–15:30", "Project deep work", "Flagship project coding off camera; launch RunPod jobs (spend cap set)", "Commits + numbers"],
    ["15:30–16:30", "Review + publish", "Watch renders (long A/B, Shorts), thumbnails, schedule uploads", "2 long + 2 Shorts scheduled"],
    ["16:30–17:15", "Social", "Post X/LinkedIn/IG/TikTok/FB from the sheet's post text; reply to every comment", "Platform columns ticked"],
    ["17:15–18:00", "Interview prep", "1 ML-systems design topic or 1 behavioral story; LC150 timed problem", "Notes"],
    ["21:00–21:40", "Tomorrow's prep", "Read both prep sheets, run the code once, watch the explainer. Stop screens 22:30", ""],
    ["Sunday 10:00–12:00", "Weekly review", "Scorecard tab; adjust next week; check DeepSeek Watch + arXiv", "Scorecard row"],
    ["Sunday 13:00–15:00", "Batch work", "Thumbnails for the week (template per module), Claude renders next week's explainers", "12 thumbnails"],
]
ws = sheet("Timetable", ["Time", "Block", "What exactly", "Output"], [20, 26, 90, 30], tt, freeze="A2")

# ------------------------------------------------------------------ Production
prod = [
    ["Night before", "Read prep sheet + script row; run `dl test <code>`; watch the explainer (explainers/out/<code>.mp4)", "You", "20 min/ep"],
    ["Explainer", "Manim film rendered from the outline's explainer lines: `dl explainer <code>` (explainers/README.md)", "Claude (batch on Sunday / night before)", "0 of your time"],
    ["Record", "Meld 1080p30, TALK (face) / CODE (screen + face circle). 9 beats: Hook, Question, Intuition [explainer plays], Math, Code, Receipt, Measure, Takeaway, Next", "You", "45–75 min/ep"],
    ["Mark done", "`dl done <code>` moves the recording into edit/<code>/ and logs it", "You", "1 min"],
    ["Edit (long)", "HyperFrames: transcribe → captions (subs.py style) + equation/title overlays (talking-head-recut) + explainer inserted at the INTUITION beat + chapter markers → render 1080p", "Claude", "≈15 min of your review"],
    ["Short", "9:16 cut of the beat named in the 'Short' column (or the vertical explainer) with burned captions, 45–60 s", "Claude", "5 min review"],
    ["Thumbnail", "Canva template per module colour: 3–5 huge words + face right third", "You (Sunday batch)", "5 min/ep"],
    ["Upload", "YouTube: title/desc/chapters from the sheet, playlist = module, end screen → next episode", "You", "10 min/ep"],
    ["Cross-post", "Same vertical file to YT Shorts, IG Reels, TikTok, FB Reels; X clip + post text; LinkedIn per Publishing tab", "You", "10 min/day"],
    ["Archive", "`archive push --purge` (dsa-content script) to hkt460s; disk is tight", "You", "weekly"],
]
sheet("Production", ["Step", "How", "Who", "Time"], [16, 110, 28, 18], prod, freeze="A2")

# ------------------------------------------------------------------ Publishing
pub = [
    ["YouTube (long)", "16:9 1080p, 10–15 min", "2/day: slot A 09:00 PT, slot B 17:00 PT (or 1/day + buffer if you prefer)", "Title from sheet; description = insight + chapters + code link + playlist link; tags: pytorch, llm, transformer, <topic>", "Playlists per module; series playlist; pin a comment with the repo link"],
    ["YouTube Shorts", "9:16, 45–60 s, burned captions", "2/day 12:00 PT", "Hook line as on-screen title; link to the long video via 'Related video'", "The ELI5/intuition or receipt beat"],
    ["Instagram Reels", "Same vertical file", "1–2/day 12:30 PT", "Hook + 3-line takeaway + 'full video on YouTube'", "Cover frame = thumbnail text"],
    ["TikTok", "Same vertical file", "1–2/day 13:00 PT", "Hook as first caption line; 3 hashtags max (#machinelearning #pytorch #llm)", "Reply to comments with video replies"],
    ["Facebook (Reels + page)", "Same vertical file", "1/day", "Same as IG", "Low effort, cross-post"],
    ["X", "Clip ≤ 2:20 or code screenshot + post", "2/day 08:30 and 18:00 PT", "x_post column; tag paper authors/repos when you reproduce them", "Reply daily to ML researchers' posts with real substance"],
    ["LinkedIn", "Native video or carousel", "3/week Tue/Thu/Sat", "linkedin column: what I built, the number, what I learned", "Module-complete recaps + project launches are the big ones"],
    ["harikanth.site", "/learn/llm track page + explainer embeds", "On each upload", "Reuse PatternExplainer component with explainers.json", "This is the page recruiters click"],
]
sheet("Publishing", ["Platform", "Format", "Cadence + time", "Caption", "Notes"], [20, 30, 40, 70, 50], pub, freeze="A2")

# ------------------------------------------------------------------ Sources
rows = []
for k, s in src.items():
    used = sorted({e["module"] for e in eps if k in e["sources"]})
    rows.append([k, s["name"], s["kind"], s["url"], ", ".join(used), len([e for e in eps if k in e["sources"]])])
rows.sort(key=lambda r: (-r[5], r[0]))
sheet("Sources", ["Key", "Source", "Type", "URL", "Modules", "Episodes using it"], [16, 60, 14, 70, 30, 10], rows, freeze="B2", wrap=False)

# ------------------------------------------------------------------ Budget
bud = [
    ["M07 GPUs/CUDA/Triton", "RTX 4090 (community)", 20, 0.34], ["M08 FlashAttention", "RTX 4090", 10, 0.34],
    ["M06E16 NanoGPT speedrun", "8x H100 SXM", 1, 8 * 2.69], ["M09 Modern architectures", "RTX 4090", 6, 0.34],
    ["M10 Distributed (real scaling)", "4x RTX 4090", 6, 4 * 0.34], ["M11 Post-training", "H100 PCIe", 25, 1.99],
    ["M12 Quantization", "RTX 4090", 12, 0.34], ["M13 Inference engines", "RTX 4090", 25, 0.34],
    ["M14 Evals/interp", "RTX 4090", 8, 0.34], ["M15 DeepSeek playbook", "H100 SXM + 4x 4090", 20, 2.69],
    ["M16 Research frontier", "H100 SXM", 40, 2.69],
    ["P1 feather-attn", "RTX 4090", 15, 0.34], ["P2 featherserve", "RTX 4090", 30, 0.34],
    ["P3 feathertrain", "8x H100 SXM", 4, 8 * 2.69], ["P4 Reasoning RL", "H100 SXM", 60, 2.69],
    ["P5 Speedrun-quality pretraining", "8x H100 SXM", 2, 8 * 2.69], ["P6 DeepSeek in miniature", "4x H100 SXM", 12, 4 * 2.69],
    ["P7 Safety research", "H100 SXM + API credits", 30, 2.69], ["P8 Low precision", "RTX 4090 + H100", 20, 1.0],
]
rows = [[a, b, h, r, f"=C{i}*D{i}", "", ""] for i, (a, b, h, r) in enumerate(bud, 2)]
ws = sheet("RunPod Budget", ["Work", "GPU", "Est. hours", "$/hr (runpod.io 2026-10-16)", "Est. $", "Actual hours", "Actual $"],
           [32, 26, 10, 14, 10, 12, 10], rows, freeze="A2")
n = ws.max_row
ws.append(["TOTAL", "", f"=SUM(C2:C{n})", "", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})"])
ws.cell(row=n + 1, column=1).font = Font(bold=True)
ws.append(["Already funded: $60 RunPod credits (bought Oct 2026). Credits applied for: see funding/GPU-CREDITS.md"])
ws.append(["Rules: set a spend limit before every pod; stop pods when idle; keep checkpoints on a network volume; prices change, re-check runpod.io/pricing."])

# ------------------------------------------------------------------ Scorecard
weeks, d0 = [], dt.date(2026, 10, 18)
last = dt.date.fromisoformat(sched[-1]["date"])
while d0 <= last + dt.timedelta(days=7):
    weeks.append([d0, 12, "", "", "", "", "", "", "", ""])
    d0 += dt.timedelta(days=7)
ws = sheet("Scorecard", ["Week ending (Sun)", "Target videos", "Videos recorded", "Long published", "Shorts published",
                         "Applications", "Outreach", "DSA episodes", "Project progress", "Notes / what to change"],
           [16, 10, 12, 12, 12, 12, 10, 12, 40, 50], weeks, freeze="B2", wrap=False)
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"

os.makedirs(f"{ROOT}/sheets", exist_ok=True)
out = f"{ROOT}/sheets/LLM-Roadmap.xlsx"
wb.save(out)
print(f"wrote {out}: {len(eps)} episodes, {len(outl)} with scripts, {len(sched)} calendar rows")
